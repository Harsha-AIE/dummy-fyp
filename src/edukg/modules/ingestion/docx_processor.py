from __future__ import annotations

from pathlib import Path

from edukg.core.exceptions import DocumentProcessingError
from edukg.core.logging_setup import get_logger
from edukg.core.models.enums import ExtractionMethod

logger = get_logger(__name__)


class DOCXProcessor:
    """Extracts text from .docx files via python-docx: paragraphs in document order
    plus table cell text, since DOCX has no native concept of 'pages'."""

    def extract(self, file_path: Path) -> str:
        try:
            import docx
        except ImportError as exc:
            raise DocumentProcessingError("python-docx is not installed") from exc

        try:
            document = docx.Document(str(file_path))
        except Exception as exc:  # noqa: BLE001
            raise DocumentProcessingError(f"Failed to open DOCX '{file_path.name}': {exc}") from exc

        parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        text = "\n".join(parts)
        logger.info("docx_extracted", file_name=file_path.name, char_count=len(text))
        return text

    def extraction_method(self) -> ExtractionMethod:
        return ExtractionMethod.NATIVE_TEXT
